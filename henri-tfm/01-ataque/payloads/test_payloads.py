"""Tests de regresión para los generadores de payload del ataque #7.

Verifican, para cada formato (PDF/DOCX/XLSX), que:
  1. El documento SANO no contiene el IBAN objetivo en ningún sitio extraíble
     (control negativo — así sabemos que el "positivo" del comprometido no es ruido).
  2. El documento COMPROMETIDO sí lo contiene, usando el mismo tipo de extracción
     ingenua que usaría un pipeline sin medidas de seguridad (sin distinguir texto
     visible de oculto).

Esto formaliza como test automatizado la verificación manual hecha durante el diseño del
payload (ver `01-ataque/README.md`, sección "Por qué cada técnica funciona").

Uso:
    ./.venv/bin/pytest -v
"""

from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from docx import Document as DocxDocument

import generar_docx
import generar_pdf
import generar_xlsx

TARGET_ACCOUNT = "ES3421000418450200051334"  # Ana Fernández Ruiz (usr_003)


# --- PDF ---

def _extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def test_pdf_sano_no_contiene_payload(tmp_path):
    out = tmp_path / "sana.pdf"
    generar_pdf.generar_sana(str(out))
    assert TARGET_ACCOUNT not in _extract_pdf_text(out)


def test_pdf_comprometido_contiene_payload(tmp_path):
    out = tmp_path / "comprometida.pdf"
    generar_pdf.generar_comprometida(str(out))
    assert TARGET_ACCOUNT in _extract_pdf_text(out)


# --- DOCX ---

def _extract_docx_text(path: Path) -> str:
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def test_docx_sano_no_contiene_payload(tmp_path):
    out = tmp_path / "sana.docx"
    generar_docx.generar_sana(str(out))
    assert TARGET_ACCOUNT not in _extract_docx_text(out)


def test_docx_comprometido_contiene_payload(tmp_path):
    out = tmp_path / "comprometida.docx"
    generar_docx.generar_comprometida(str(out))
    assert TARGET_ACCOUNT in _extract_docx_text(out)


def test_docx_comprometido_usa_atributo_hidden(tmp_path):
    """Confirma que la técnica es específicamente `run.font.hidden`, no solo el color."""
    out = tmp_path / "comprometida.docx"
    generar_docx.generar_comprometida(str(out))
    doc = DocxDocument(str(out))
    hidden_runs = [
        run for p in doc.paragraphs for run in p.runs
        if run.font.hidden and TARGET_ACCOUNT in run.text
    ]
    assert hidden_runs, "El payload debería viajar en un run con font.hidden=True"


# --- XLSX ---

def _extract_xlsx_payload_locations(path: Path) -> tuple[bool, bool]:
    """Devuelve (encontrado_en_celda_oculta, encontrado_en_comentario)."""
    wb = load_workbook(str(path))
    ws = wb.active

    found_hidden_cell = False
    for row_idx, row in enumerate(ws.iter_rows(), start=1):
        is_hidden_row = ws.row_dimensions[row_idx].hidden
        for cell in row:
            if cell.value and TARGET_ACCOUNT in str(cell.value) and is_hidden_row:
                found_hidden_cell = True

    found_comment = any(
        cell.comment and TARGET_ACCOUNT in cell.comment.text
        for row in ws.iter_rows()
        for cell in row
    )
    return found_hidden_cell, found_comment


def test_xlsx_sano_no_contiene_payload(tmp_path):
    out = tmp_path / "sano.xlsx"
    generar_xlsx.generar_sano(str(out))
    hidden_cell, comment = _extract_xlsx_payload_locations(out)
    assert not hidden_cell and not comment


def test_xlsx_comprometido_contiene_payload_en_fila_oculta(tmp_path):
    out = tmp_path / "comprometido.xlsx"
    generar_xlsx.generar_comprometido(str(out))
    hidden_cell, _ = _extract_xlsx_payload_locations(out)
    assert hidden_cell


def test_xlsx_comprometido_contiene_payload_en_comentario(tmp_path):
    out = tmp_path / "comprometido.xlsx"
    generar_xlsx.generar_comprometido(str(out))
    _, comment = _extract_xlsx_payload_locations(out)
    assert comment
