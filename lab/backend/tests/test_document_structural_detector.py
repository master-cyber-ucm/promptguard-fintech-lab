"""Tests de `src.core.document_structural_detector` — ataque #7, defensa (A) parcial.

Genera documentos mínimos con las mismas técnicas de ocultación usadas en los payloads del
ataque (ver henri-tfm/01-ataque/anatomia-payload.md) y verifica que se detectan como firma
conocida — y que los documentos sanos equivalentes no disparan ningún falso positivo.
"""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.comments import Comment
from reportlab.lib.colors import black, white
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from src.core.document_structural_detector import detect_hiding_techniques


# --- PDF ---

def _pdf_bytes(white_text=False, tiny_font=False, outside_page=False) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    c.setFillColor(black)
    c.setFont("Helvetica", 12)
    c.drawString(60, height - 60, "Documento visible de prueba")
    if white_text:
        c.setFillColor(white)
        c.drawString(60, height - 100, "Texto blanco sobre blanco")
        c.setFillColor(black)
    if tiny_font:
        c.setFont("Helvetica", 1)
        c.drawString(60, height - 120, "Texto en fuente diminuta")
        c.setFont("Helvetica", 12)
    if outside_page:
        c.drawString(60, -50, "Texto fuera de la pagina")
    c.save()
    return buf.getvalue()


def test_pdf_sano_no_detecta_nada():
    assert detect_hiding_techniques("doc.pdf", _pdf_bytes()) == []


def test_pdf_detecta_texto_blanco():
    findings = detect_hiding_techniques("doc.pdf", _pdf_bytes(white_text=True))
    assert "pdf_white_text" in findings


def test_pdf_detecta_fuente_diminuta():
    findings = detect_hiding_techniques("doc.pdf", _pdf_bytes(tiny_font=True))
    assert "pdf_tiny_font" in findings


def test_pdf_detecta_texto_fuera_de_pagina():
    findings = detect_hiding_techniques("doc.pdf", _pdf_bytes(outside_page=True))
    assert "pdf_text_outside_page" in findings


# --- DOCX ---

def _docx_bytes(hidden=False) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Documento visible de prueba")
    if hidden:
        run = doc.add_paragraph().add_run("Texto oculto")
        run.font.hidden = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(1)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_sano_no_detecta_nada():
    assert detect_hiding_techniques("doc.docx", _docx_bytes()) == []


def test_docx_detecta_run_oculto():
    findings = detect_hiding_techniques("doc.docx", _docx_bytes(hidden=True))
    assert "docx_hidden_run" in findings


# --- XLSX ---

def _xlsx_bytes(hidden_row=False, comment=False) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Concepto", "Importe"])
    ws.append(["Gasto de prueba", -10.0])
    if hidden_row:
        idx = ws.max_row + 1
        ws.cell(row=idx, column=1, value="fila oculta")
        ws.row_dimensions[idx].hidden = True
    if comment:
        ws.cell(row=2, column=1).comment = Comment("nota oculta", "Sistema")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_sano_no_detecta_nada():
    assert detect_hiding_techniques("doc.xlsx", _xlsx_bytes()) == []


def test_xlsx_detecta_fila_oculta():
    findings = detect_hiding_techniques("doc.xlsx", _xlsx_bytes(hidden_row=True))
    assert "xlsx_hidden_row" in findings


def test_xlsx_detecta_comentario():
    findings = detect_hiding_techniques("doc.xlsx", _xlsx_bytes(comment=True))
    assert "xlsx_cell_comment" in findings


def test_formato_no_soportado_devuelve_lista_vacia():
    assert detect_hiding_techniques("archivo.txt", b"contenido") == []
