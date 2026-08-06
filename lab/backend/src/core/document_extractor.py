"""Extracción de texto de documentos adjuntos — MODO VULNERABLE.

Ataque #7 (OWASP LLM01:2025 · MITRE ATLAS AML.T0051.001) — Prompt Injection Indirecta vía
Documento.

Extrae texto de PDF/DOCX/XLSX de forma deliberadamente ingenua: no distingue contenido visible
de oculto (color de fuente, tamaño, filas ocultas de hoja de cálculo, comentarios de celda,
atributo "hidden" de Word), no sanitiza, no separa semánticamente "documento" de "instrucción".
Esto reproduce el patrón de fallo ya documentado en
docs/ataques/LLM01-prompt-injection/indirecta-documento/04-analisis-tecnico.md: el pipeline trata
el texto extraído como dato confiable, indistinguible del resto del prompt.

La Fase 2 (defensa) del TFM tratará este texto como input NO confiable, pasándolo por el mismo
pipeline que el Input Sanitizer aplicaría al mensaje de chat directo — ver
henri-tfm/02-defensa/README.md. Este módulo no implementa ninguna defensa a propósito.
"""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}


class UnsupportedDocumentError(ValueError):
    """El documento adjunto no tiene un formato soportado por el pipeline de extracción."""


def extract_text(filename: str, content: bytes) -> str:
    """Extrae todo el texto recuperable de un documento, sin filtrar por visibilidad.

    Lanza UnsupportedDocumentError si la extensión no está en SUPPORTED_EXTENSIONS.
    """
    ext = _extension(filename)
    if ext == ".pdf":
        return _extract_pdf(content)
    if ext == ".docx":
        return _extract_docx(content)
    if ext == ".xlsx":
        return _extract_xlsx(content)
    raise UnsupportedDocumentError(
        f"Formato no soportado: '{ext or filename}'. Soportados: {sorted(SUPPORTED_EXTENSIONS)}"
    )


def _extension(filename: str) -> str:
    if "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[-1].lower()


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    # NOTA: python-docx `paragraph.text` concatena todos los runs, incluidos los que tienen
    # `run.font.hidden = True` (atributo w:vanish de Word). No filtra por visibilidad.
    doc = DocxDocument(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_xlsx(content: bytes) -> str:
    # NOTA: se recorren TODAS las filas (incluidas las marcadas como `hidden`) y también los
    # comentarios de celda — ninguna de las dos superficies se filtra.
    wb = load_workbook(io.BytesIO(content))
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    lines.append(str(cell.value))
                if cell.comment is not None:
                    lines.append(cell.comment.text)
    return "\n".join(lines)
