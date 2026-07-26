"""Detección estructural de técnicas de ocultación conocidas — Fase 2, ataque #7.

⚠️ ESTO ES UNA BASE DE FIRMAS, NO UNA DEFENSA COMPLETA — análogo a un antivirus de firmas.
Cubre ÚNICAMENTE las técnicas de ocultación catalogadas en
`henri-tfm/01-ataque/anatomia-payload.md` a la fecha de este trabajo. El análisis completo de por
qué esto NO sustituye a `document_sanitizer.py` (la capa base, agnóstica a la técnica de
ocultación) está en `henri-tfm/02-defensa/README.md` §"Análisis de viabilidad de (A)": el catálogo
de técnicas de esteganografía de texto en documentos es amplio y sigue creciendo (caracteres
Unicode invisibles, homoglifos, capas de contenido opcional de PDF, objetos incrustados...) —
ninguna de esas técnicas adicionales se detecta aquí.

Este módulo debe tratarse como un catálogo VIVO, igual que una base de firmas de antivirus: cada
vez que se documente una técnica de ocultación nueva (nuevo hallazgo de red-teaming, nueva
variante de ataque), el proceso es:
  1. Añadir un test de regresión que la reproduzca.
  2. Añadir la detección correspondiente en la función del formato afectado.
  3. Registrar la entrada en el CHANGELOG de abajo (versión, fecha, motivo).

No se ha medido su coste en tiempo de ejecución de forma abstracta — ver
`henri-tfm/02-defensa/benchmark_structural_detector.py` para la medición real.

CHANGELOG de firmas
--------------------
- v1 (2026-07-26) — 5 técnicas iniciales, las mismas caracterizadas en la Fase 1 de este ataque:
  · PDF: color de texto = blanco puro (RGB 1,1,1 / gris 1.0) — NO detecta "color similar al
    fondo" en general, solo blanco puro, que es la técnica exacta usada en la Fase 1. Detectar
    "color parecido al fondo" de forma genérica requeriría conocer el color de fondo real de la
    página, un problema de renderizado fuera de alcance de esta versión.
  · PDF: tamaño de fuente < 2pt.
  · PDF: texto con coordenada Y fuera del alto de página (incluye "fuera de viewport").
  · DOCX: atributo `run.font.hidden` (`w:vanish`).
  · XLSX: fila/columna con `hidden=True`, o celda con comentario adjunto.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from .document_extractor import _extension

PDF_HIDDEN_FONT_SIZE_THRESHOLD = 2.0


def detect_hiding_techniques(filename: str, content: bytes) -> list[str]:
    """Devuelve la lista de técnicas de ocultación CONOCIDAS (v1, ver CHANGELOG) encontradas en
    el documento. Lista vacía si no se encuentra ninguna — no implica que el documento esté
    limpio, solo que no coincide con ninguna firma catalogada (ver advertencia del módulo).
    """
    ext = _extension(filename)
    if ext == ".pdf":
        return _detect_pdf(content)
    if ext == ".docx":
        return _detect_docx(content)
    if ext == ".xlsx":
        return _detect_xlsx(content)
    return []


def _detect_pdf(content: bytes) -> list[str]:
    findings: list[str] = []
    reader = PdfReader(io.BytesIO(content))

    state: dict[str, Any] = {"fill_color": None}

    def _op_before(operator: bytes, operands: list, cm, tm) -> None:
        if operator in (b"rg", b"sc", b"scn"):
            state["fill_color"] = list(operands)
        elif operator == b"g":
            v = operands[0] if operands else None
            state["fill_color"] = [v, v, v] if v is not None else None

    for page in reader.pages:
        page_height = float(page.mediabox.height)

        def _visitor(text: str, cm, tm, font_dict, font_size) -> None:
            if not text.strip():
                return
            color = state["fill_color"]
            if color is not None and all(_is_white(c) for c in color):
                if "pdf_white_text" not in findings:
                    findings.append("pdf_white_text")
            if font_size is not None and font_size > 0 and font_size < PDF_HIDDEN_FONT_SIZE_THRESHOLD:
                if "pdf_tiny_font" not in findings:
                    findings.append("pdf_tiny_font")
            y = tm[5] if tm else None
            if y is not None and (y < 0 or y > page_height):
                if "pdf_text_outside_page" not in findings:
                    findings.append("pdf_text_outside_page")

        page.extract_text(visitor_operand_before=_op_before, visitor_text=_visitor)

    return findings


def _is_white(value: float, tolerance: float = 0.02) -> bool:
    return abs(value - 1.0) <= tolerance


def _detect_docx(content: bytes) -> list[str]:
    doc = DocxDocument(io.BytesIO(content))
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.hidden:
                return ["docx_hidden_run"]
    return []


def _detect_xlsx(content: bytes) -> list[str]:
    findings: list[str] = []
    wb = load_workbook(io.BytesIO(content))
    for ws in wb.worksheets:
        for dim in ws.row_dimensions.values():
            if dim.hidden:
                findings.append("xlsx_hidden_row")
                break
        for dim in ws.column_dimensions.values():
            if dim.hidden:
                findings.append("xlsx_hidden_column")
                break
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    findings.append("xlsx_cell_comment")
                    break
            if "xlsx_cell_comment" in findings:
                break
    # dedupe preservando orden
    seen: set[str] = set()
    return [f for f in findings if not (f in seen or seen.add(f))]
